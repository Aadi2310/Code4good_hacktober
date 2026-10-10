"""Small dependency-backed smoke check for the P1 CSV and output path."""
import hashlib
import io
import tempfile
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from vyom.intake import detect
from vyom.models import RecordSource
from vyom.tabular import run
from vyom.outputs.export import csv_bytes, xlsx_bytes
from openpyxl import load_workbook
from vyom.api.main import app

SAMPLE = """Invoice No,Invoice Date,Supplier,Description,HSN,Qty,Unit Price,Taxable Value,CGST,SGST,Invoice Total
INV-1,2026-10-01,Acme,=Widget,0123,2,50,100,9,9,118
"""

with tempfile.TemporaryDirectory(prefix="vyom-smoke-", dir=Path(__file__).resolve().parents[1]) as directory:
    path = Path(directory) / "sample.csv"
    path.write_text(SAMPLE, encoding="utf-8", newline="")
    detected = detect(path)
    source = RecordSource(filename="sample.csv", sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                         detected_type=detected.mime, input_kind=detected.kind)
    output = run(path, source)
    assert len(output.records) == 1, f"expected 1 invoice; got {len(output.records)}"
    assert len(output.table_rows) == 1, f"expected 1 line row; got {len(output.table_rows)}"
    assert output.records[0].invoice["invoice_no"].value == "INV-1"
    assert output.records[0].line_items[0]["hsn_sac"].value == "0123"
    csv = csv_bytes(output.records)
    assert csv.startswith(b"\xef\xbb\xbf") and b"\r\n" in csv and b"'=Widget" in csv
    workbook = xlsx_bytes(output.records)
    assert workbook.startswith(b"PK")
    book = load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    assert book.sheetnames == ["Invoices", "LineItems", "Validation", "Corrections", "Audit"]
    assert app.openapi_url == "/api/v1/openapi.json"
    print("P1 smoke passed: CSV intake, 1 invoice/line, identifier preservation, CSV BOM/CRLF, XLSX and FastAPI contract")
