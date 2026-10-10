import csv,io,json
from datetime import datetime,timezone,date
from decimal import Decimal,InvalidOperation
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
from ..models import FLAT_COLUMNS

def exportable(records,scope="exportable"):
    if scope=="all":return records
    return [r for r in records if r.review.state!="rejected" and (r.validation.status=="valid" or r.review.state=="approved")]
def json_bytes(records):
    data={"export_version":"1.0","exported_at":datetime.now(timezone.utc).isoformat(),"documents":[r.model_dump(mode="json") for r in records]}
    return json.dumps(data,ensure_ascii=False,indent=2).encode("utf-8")
def _safe(v):
    if v is None:return ""
    s=str(v)
    if s[:1] in ("=","+","@","\t","\r") or (s.startswith("-") and not _number(s)):return "'"+s
    return s
def _number(s):
    try:Decimal(s);return True
    except InvalidOperation:return False
def _xlsx_value(key,value):
    if value is None:return ""
    text=str(value)
    if key in {"invoice_date","due_date"}:
        try:return date.fromisoformat(text[:10])
        except ValueError:pass
    if key in {"subtotal","discount","taxable_value","cgst","sgst","igst","cess","total_tax","round_off","total_amount","invoice_total","unit_price","line_total","quantity","tax_rate"} and _number(text):
        return Decimal(text)
    return _safe(value)
def _flat(records):
    result=[]
    for rec in records:
        header={k:(f.value or "") for k,f in rec.invoice.items()}
        lines=rec.line_items or [{}]
        for line in lines:
            row={k:"" for k in FLAT_COLUMNS};row.update(header)
            row.update({k:(f.value or "") for k,f in line.items()})
            row["invoice_total"]=header.get("total_amount",""); row["extra"]=""
            result.append(row)
    return result
def csv_bytes(records):
    s=io.StringIO(newline="");w=csv.DictWriter(s,fieldnames=FLAT_COLUMNS,extrasaction="ignore",lineterminator="\r\n");w.writeheader()
    for row in _flat(records):w.writerow({k:_safe(v) for k,v in row.items()})
    return b"\xef\xbb\xbf"+s.getvalue().encode("utf-8")
def xlsx_bytes(records):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active;ws.title="Invoices"
    header_fields=sorted({k for r in records for k in r.invoice})
    cols=header_fields+["validation_status","score","review_state","overall_confidence"]
    _sheet(ws,cols)
    for r in records:
        ws.append([_xlsx_value(k,r.invoice.get(k).value if r.invoice.get(k) else "") for k in header_fields]+[r.validation.status,r.validation.score,r.review.state,r.quality.overall_confidence])
        _color(ws,ws.max_row,r)
        for column,k in enumerate(header_fields,1):
            if k in {"subtotal","discount","taxable_value","cgst","sgst","igst","cess","total_tax","round_off","total_amount"}:ws.cell(ws.max_row,column).number_format="#,##0.00"
            if k in {"invoice_date","due_date"}:ws.cell(ws.max_row,column).number_format="yyyy-mm-dd"
    line=wb.create_sheet("LineItems");_sheet(line,FLAT_COLUMNS)
    for row in _flat(records):
        line.append([_xlsx_value(k,row.get(k)) for k in FLAT_COLUMNS])
        for column,k in enumerate(FLAT_COLUMNS,1):
            if k in {"unit_price","discount","taxable_value","cgst","sgst","igst","cess","line_total","invoice_total"}:line.cell(line.max_row,column).number_format="#,##0.00"
            if k in {"invoice_date","due_date"}:line.cell(line.max_row,column).number_format="yyyy-mm-dd"
    checks=wb.create_sheet("Validation");_sheet(checks,["record_id","code","severity","passed","field_paths","message","expected","actual","difference"])
    for r in records:
        for c in r.validation.checks:checks.append([r.record_id,c.code,c.severity,c.passed,", ".join(c.field_paths),_safe(c.message),c.expected,c.actual,c.difference])
    corr=wb.create_sheet("Corrections");_sheet(corr,["record_id","field_path","original","corrected","method","reason"])
    for r in records:
        for c in r.validation.corrections:corr.append([r.record_id,c.field_path,c.original,c.corrected,c.method,c.reason])
    audit=wb.create_sheet("Audit");_sheet(audit,["record_id","ts","actor","action","field_path","old","new"])
    for r in records:
        for e in r.review.history:audit.append([r.record_id,e.ts,e.actor,e.action,e.field_path,e.old,e.new])
    b=io.BytesIO();wb.save(b);return b.getvalue()
def _sheet(ws,headers):
    ws.append(headers);ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True)
    for column in ws.columns:
        ws.column_dimensions[get_column_letter(column[0].column)].width=min(48,max(12,max((len(str(c.value or "")) for c in column),default=0)+2))
def _color(ws,row,rec):
    color="FCE4D6" if rec.validation.status in ("needs_review","failed") else "FFF2CC" if rec.validation.status=="valid_with_warnings" else None
    if color:
        for c in ws[row]:c.fill=PatternFill("solid",fgColor=color)
