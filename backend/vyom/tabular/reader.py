"""Deterministic CSV/XLSX invoice ingestion (layouts A, B and C)."""
from __future__ import annotations
import csv, io, re, uuid, zipfile, xml.etree.ElementTree as ET, posixpath
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from decimal import ROUND_HALF_UP
from charset_normalizer import from_bytes
from rapidfuzz import fuzz
from ..config import MAX_XLSX_ROWS, MAX_XLSX_SHEETS, MAX_UNCOMPRESSED_BYTES
from ..models import HEADER_FIELDS, LINE_FIELDS, FLAT_COLUMNS, FV, InvoiceRecord, RecordSource, TabularOutput, Check
from ..domain import normalize_field, parse_money
from ..validation import validate
from ..errors import PipelineError

ALIASES = {
 "invoice_no":["invoice no","inv no","bill no","bill number","tax invoice no","invoice number","challan no","voucher no","document no","doc no"],
 "invoice_date":["invoice date","inv date","bill date","date","doc date","dinank","tarikh"],
 "due_date":["due date","payment due date"], "document_type":["type","document type","voucher type","invoice type"],
 "place_of_supply":["place of supply","pos","state of supply"], "purchase_order_no":["purchase order no","po no","po number"],
 "reverse_charge":["reverse charge","rcm"], "supplier_name":["supplier","seller","vendor","party name","from","supplier name"],
 "supplier_gstin":["supplier gstin","seller gstin","vendor gstin","gstin of supplier","gstin"],
 "buyer_name":["buyer","customer","recipient","receiver name","bill to","buyer name","customer name"],
 "buyer_gstin":["buyer gstin","customer gstin","recipient gstin","gstin of recipient"],
 "line_no":["sr no","s no","sl no","serial","sr"], "description":["description","particulars","item","item name","product","goods","description of goods"],
 "hsn_sac":["hsn","hsn sac","hsn code","sac","hsn sac code"], "quantity":["qty","quantity","nos","pcs","units"],
 "unit":["uom","unit","unit of measure"], "unit_price":["unit price","rate per unit","price","mrp","basic rate"],
 "discount":["discount","disc","discount amount"], "taxable_value":["taxable value","taxable amt","taxable amount","assessable value"],
 "cgst_rate":["cgst %","cgst rate","c g s t rate"], "cgst":["cgst amt","cgst amount","cgst"],
 "sgst_rate":["sgst %","sgst rate","utgst %","utgst rate"], "sgst":["sgst amt","sgst amount","sgst","utgst amount"],
 "igst_rate":["igst %","igst rate"], "igst":["igst amt","igst amount","igst"],
 "cess_rate":["cess %","cess rate"], "cess":["cess amt","cess amount","cess"],
 "line_total":["line total","amount","amt","value"], "subtotal":["subtotal","sub total"],
 "total_tax":["total tax","total gst","tax amount"], "round_off":["round off","rounding"],
 "invoice_total":["invoice value","invoice total","grand total","total invoice value","net payable","amount payable"],
 "amount_in_words":["amount in words","rupees in words"]
}
RATE_KEYS={"cgst_rate","sgst_rate","igst_rate","cess_rate"}
NUMERIC_KEYS={"quantity","unit_price","discount","taxable_value","cgst","sgst","igst","cess","line_total","subtotal","total_tax","round_off","invoice_total",*RATE_KEYS}
TOTAL_RE=re.compile(r"^(sub\s*total|total|grand\s*total|round\s*off|rounded|net\s*payable|amount\s*payable)\b",re.I)
GST_RATES={Decimal(x) for x in ("0","0.05","0.1","0.25","0.5","1","1.5","2.5","3","4","5","6","7.5","9","12","14","18","28")}

def _clean(value):
    if value is None:return ""
    if isinstance(value,bool):return str(value)
    if isinstance(value,(datetime,date)):return value.isoformat()
    if isinstance(value,float) and value.is_integer():value=int(value)
    text=str(value).replace("\u00a0"," ").replace("\u200b","").strip()
    if text.startswith("'"):text=text[1:]
    return text.translate(str.maketrans("०१२३४५६७८९","0123456789"))
def _norm(value):return re.sub(r"[^a-z0-9%]+"," ",_clean(value).lower()).strip()
def _unique_headers(headers):
    used={};out=[]
    for i,h in enumerate(headers):
        base=_clean(h) or f"col_{i+1}";used[base]=used.get(base,0)+1
        out.append(base if used[base]==1 else f"{base}_{used[base]}")
    return out
def _header_score(row):
    matches=0
    for cell in row:
        n=_norm(cell)
        if n and any(n==_norm(a) or fuzz.ratio(n,_norm(a))>=88 for als in ALIASES.values() for a in als):matches+=1
    return matches
def _find_header(rows):
    best_i=-1;best=0
    for i,row in enumerate(rows[:200]):
        score=_header_score(row)
        if score>best:best_i,best=i,score
    for i in range(min(199,len(rows)-1)):
        upper,lower=rows[i],rows[i+1]
        lower_components=sum(bool(_clean(x)) and _norm(x) in {"rate","amount","amt","percentage","percent"} for x in lower)
        if _header_score(upper)<2 or (_header_score(lower)<2 and lower_components<2):continue
        joined=[];previous=""
        for j in range(max(len(upper),len(lower))):
            a=_clean(upper[j]) if j<len(upper) else "";b=_clean(lower[j]) if j<len(lower) else ""
            if not a:a=previous
            elif a:previous=a
            joined.append(f"{a} {b}".strip() if a and b and _norm(a)!=_norm(b) else b or a)
        if _header_score(joined)>=4:return i+1,joined
    if best_i<0 or best<3:return None,None
    # Join two-tier headers, e.g. a merged CGST cell over Rate and Amount.
    if best_i+1<len(rows) and _header_score(rows[best_i+1])>=2:
        lower=rows[best_i+1];upper=rows[best_i]; joined=[]
        for j in range(max(len(upper),len(lower))):
            a=_clean(upper[j]) if j<len(upper) else "";b=_clean(lower[j]) if j<len(lower) else ""
            joined.append(f"{a} {b}".strip() if a and b and _norm(a)!=_norm(b) else b or a)
        if _header_score(joined)>=4:return best_i+1,joined
    return best_i,rows[best_i]
def _two_level_header(rows,lower_index):
    if lower_index<=0 or lower_index>=len(rows):return False
    return _header_score(rows[lower_index-1])>=2 and sum(_norm(x) in {"rate","amount","amt","percentage","percent"} for x in rows[lower_index] if _clean(x))>=2

def _decode_csv(raw):
    if raw.startswith(b"\xef\xbb\xbf"):return raw.decode("utf-8-sig"),"utf-8-sig"
    if raw.startswith((b"\xff\xfe",b"\xfe\xff")):return raw.decode("utf-16"),"utf-16"
    detected=from_bytes(raw).best()
    if detected:return str(detected),detected.encoding or "unknown"
    try:return raw.decode("utf-8"),"utf-8"
    except UnicodeDecodeError:return raw.decode("cp1252"),"cp1252"
def _choose_delimiter(text):
    sample=text[:65536]
    def score(d):
        try:counts=[len(r) for r in list(csv.reader(io.StringIO(sample),delimiter=d))[:50] if r and any(x.strip() for x in r)]
        except csv.Error:return (0,0)
        if not counts:return (0,0)
        mode=max(set(counts),key=counts.count)
        return (sum(n==mode for n in counts),mode) if mode>1 else (0,0)
    try:guessed=csv.Sniffer().sniff(sample,delimiters=",;\t|").delimiter
    except csv.Error:guessed=None
    if guessed and score(guessed)[0]>=2:return guessed
    return max(",;\t|",key=score)
def _read_csv(path):
    raw=Path(path).read_bytes()
    text,encoding=_decode_csv(raw);delimiter=_choose_delimiter(text)
    try:
        reader=csv.reader(io.StringIO(text,newline=""),delimiter=delimiter,strict=False);parsed=[];previous_line=0
        for record in reader:
            end=reader.line_num;parsed.append((previous_line+1,record));previous_line=end
    except csv.Error as e:raise PipelineError("CORRUPT_FILE","CSV could not be parsed") from e
    rows=[record for _,record in parsed]
    if not rows:return [],[],encoding,[],[]
    header_i,headers=_find_header(rows)
    found_header=headers is not None
    warnings=[]
    if headers is None:
        headers=[f"col_{i+1}" for i in range(max(map(len,rows)))];body=rows;warnings.append(Check(code="HEADER_NOT_FOUND",severity="warning",passed=True,message="Generated column names because no recognizable header was found"))
    else:body=rows[header_i+1:]
    preamble=rows[:header_i-1 if _two_level_header(rows,header_i) else header_i] if found_header else []
    headers=_unique_headers(headers);width=len(headers);padded=[]
    body_origins=parsed[header_i+1:] if found_header else parsed
    for physical_line,row in body_origins:
        if not any(_clean(x) for x in row):padded.append((physical_line,[""]*width));continue
        if len(row)>width:
            old=width;headers.extend(f"extra_{i+1}" for i in range(len(row)-old));width=len(row)
            warnings.append(Check(code="RAGGED_ROWS",severity="warning",passed=True,message="Rows wider than the detected header were retained in extra columns"))
        padded.append((physical_line,[_clean(x) for x in row]+[""]*max(0,width-len(row))))
    if len(padded)>MAX_XLSX_ROWS:raise PipelineError("TOO_MANY_ROWS","CSV exceeds the configured row limit")
    return headers,_trim_blank_run(padded),encoding,warnings,preamble

def _trim_blank_run(rows):
    run=0;end=len(rows)
    for i,(_,row) in enumerate(rows):
        if not any(_clean(x) for x in row):
            run+=1
            if run>=3:end=i-2;break
        else:run=0
    return rows[:end]

def _read_xlsx(path):
    try:import openpyxl
    except ImportError as e:raise PipelineError("DEPENDENCY_MISSING","Install openpyxl to read XLSX") from e
    try:
        with zipfile.ZipFile(path) as z:
            infos=z.infolist()
            if sum(x.file_size for x in infos)>MAX_UNCOMPRESSED_BYTES:raise PipelineError("ARCHIVE_TOO_LARGE","XLSX expands beyond the configured safe size")
            if any(x.filename.startswith("/") or ".." in Path(x.filename).parts for x in infos):raise PipelineError("ARCHIVE_UNSAFE","XLSX contains an unsafe archive path")
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
        formulas=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    except PipelineError:raise
    except Exception as e:raise PipelineError("CORRUPT_FILE","Unable to open XLSX") from e
    if len(wb.worksheets)>MAX_XLSX_SHEETS:
        wb.close();formulas.close();raise PipelineError("TOO_MANY_SHEETS","XLSX exceeds the configured sheet limit")
    outputs=[];global_warnings=[];metadata=_xlsx_metadata(path)
    if any(s.sheet_state!="visible" for s in wb.worksheets):global_warnings.append(Check(code="HIDDEN_CONTENT_IGNORED",severity="warning",passed=True,message="Hidden sheets were ignored"))
    for ws in wb.worksheets:
        if ws.sheet_state!="visible":continue
        fws=formulas[ws.title];rows=[];frows=fws.iter_rows(values_only=True);meta=metadata.get(ws.title,{})
        for ri,row in enumerate(ws.iter_rows(values_only=True),1):
            if ri>MAX_XLSX_ROWS:raise PipelineError("TOO_MANY_ROWS",f"Sheet {ws.title} exceeds the row limit")
            formula_row=next(frows,())
            vals=[]
            for ci,value in enumerate(row,1):
                raw_formula=formula_row[ci-1] if ci-1<len(formula_row) else None
                if isinstance(raw_formula,str) and raw_formula.startswith("=") and value is None:
                    global_warnings.append(Check(code="FORMULAS_NOT_CALCULATED",severity="warning",passed=True,message=f"Uncalculated formula encountered in {ws.title}!{openpyxl.utils.get_column_letter(ci)}{ri}"))
                vals.append(_clean(value))
            rows.append((ri,vals))
        try:
            from openpyxl.utils.cell import range_boundaries
            rowmap={ri:vals for ri,vals in rows};hidden_rows=meta.get("hidden_rows",set());hidden_columns=meta.get("hidden_columns",set())
            for merge in meta.get("merged",[]):
                minc,minr,maxc,maxr=range_boundaries(merge);anchor=rowmap.get(minr,[]);value=anchor[minc-1] if minc-1<len(anchor) else ""
                for rr in range(minr,maxr+1):
                    target=rowmap.setdefault(rr,[])
                    if len(target)<maxc:target.extend([""]*(maxc-len(target)))
                    for cc in range(minc,maxc+1):
                        if not target[cc-1]:target[cc-1]=value
            if hidden_rows or hidden_columns:
                global_warnings.append(Check(code="HIDDEN_CONTENT_IGNORED",severity="warning",passed=True,message=f"Ignored {len(hidden_rows)} hidden rows and {len(hidden_columns)} hidden columns in {ws.title}"))
            rows=[(ri,[_clean(v) for ci,v in enumerate(vals,1) if ci not in hidden_columns]) for ri,vals in sorted(rowmap.items()) if ri not in hidden_rows]
        except Exception:pass
        nonempty=[(ri,row) for ri,row in rows if any(row)]
        if len(nonempty)<2:continue
        header_i,headers=_find_header([row for _,row in nonempty])
        if headers is None:
            width=max(map(lambda pair:len(pair[1]),rows));headers=[f"col_{i+1}" for i in range(width)];body=rows
            preamble=[];actual_header_row=0
            global_warnings.append(Check(code="HEADER_NOT_FOUND",severity="warning",passed=True,message=f"Generated headers on sheet {ws.title}"))
        else:
            actual_header_row=nonempty[header_i][0]
            start_row=nonempty[header_i-1][0] if _two_level_header([row for _,row in nonempty],header_i) else actual_header_row
            preamble=[row for ri,row in rows if ri<start_row]
            # Consume all physical rows after the header; preserve the actual worksheet row for provenance.
            body=[(ri,row) for ri,row in rows if ri>actual_header_row]
        outputs.append((ws.title,_unique_headers(headers),_trim_blank_run(body),actual_header_row if headers else 0,preamble))
    wb.close();formulas.close();return outputs,global_warnings

def _xlsx_metadata(path):
    ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","rel":"http://schemas.openxmlformats.org/package/2006/relationships"}
    try:
        with zipfile.ZipFile(path) as archive:
            workbook=ET.fromstring(archive.read("xl/workbook.xml"));rels=ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relmap={r.attrib["Id"]:r.attrib["Target"] for r in rels.findall("rel:Relationship",ns)};meta={}
            for sheet in workbook.findall("m:sheets/m:sheet",ns):
                title=sheet.attrib["name"];rid=sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id");target=relmap.get(rid,"")
                xmlpath=posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl",target))
                root=ET.fromstring(archive.read(xmlpath));hidden_rows={int(r.attrib["r"]) for r in root.findall(".//m:row",ns) if r.attrib.get("hidden")=="1"};hidden_columns=set()
                for col in root.findall(".//m:col",ns):
                    if col.attrib.get("hidden")=="1":hidden_columns.update(range(int(col.attrib["min"]),int(col.attrib["max"])+1))
                meta[title]={"hidden_rows":hidden_rows,"hidden_columns":hidden_columns,"merged":[x.attrib["ref"] for x in root.findall(".//m:mergeCell",ns)]}
            return meta
    except Exception:return {}

def _mapping(headers,rows):
    result={};used=set()
    for i,h in enumerate(headers):
        n=_norm(h);best=None;score=0;method=""
        for key,aliases in ALIASES.items():
            if n in {_norm(a) for a in aliases}:best=key;score=1.;method="exact";break
            q=max((fuzz.token_set_ratio(n,_norm(a))/100 for a in aliases),default=0)
            if q>=.88 and q>score:best=key;score=q;method="fuzzy"
        if best and best not in used:result[i]=(best,method,score);used.add(best)
    # Content-based typing for otherwise unmapped columns.
    for i,h in enumerate(headers):
        if i in result:continue
        vals=[_clean(r[i]) for r in rows[:200] if i<len(r) and _clean(r[i])]
        if not vals:continue
        gst=sum(bool(re.fullmatch(r"\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]",v.upper())) for v in vals)/len(vals)
        date_ratio=sum(_date_value(v)[0] is not None for v in vals)/len(vals)
        ident=sum(bool(re.fullmatch(r"\d{4}(?:\d{2}){0,2}",v)) for v in vals)/len(vals)
        rate=sum(_decimal(v) in GST_RATES for v in vals)/len(vals)
        target=None
        if gst>=.8:target="buyer_gstin"
        elif date_ratio>=.8:target="invoice_date"
        elif ident>=.8:target="hsn_sac"
        elif rate>=.8:target="tax_rate"
        if target and target not in used:result[i]=(target,"content",.7 if target=="hsn_sac" else .8);used.add(target)
    return result
def _decimal(v):
    try:return Decimal(str(v).replace(",","").replace("%","").strip())
    except (InvalidOperation,ValueError):return None
def _date_value(value):
    if not value:return None,[]
    try:return datetime.fromisoformat(str(value).strip().replace("Z","+00:00")).date().isoformat(),[]
    except ValueError:pass
    for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%d.%m.%Y","%d/%m/%y"):
        try:return datetime.strptime(str(value).strip(),fmt).date().isoformat(),[]
        except ValueError:pass
    return None,["invalid_date"]
def _valmap(row,mapping):return {key:_clean(row[i]) if i<len(row) else "" for i,(key,_,_) in mapping.items()}
def _mk_fv(name,raw,source_kind,ref,confidence=.97):
    if raw is None or raw=="":return FV()
    normal=raw
    if name in RATE_KEYS or name=="tax_rate":
        d=_decimal(raw)
        if d is None:return FV(raw=raw,confidence=confidence,status="needs_review",sources=[source_kind],source_ref=ref,evidence_text=raw)
        if d<=1 and d*100 in GST_RATES:d*=100
        normal=format(d.normalize(),"f")
        return FV(value=normal,raw=raw,confidence=confidence,status="accepted",sources=[source_kind],source_ref=ref,evidence_text=raw)
    if name=="quantity":
        d=_decimal(raw)
        if d is None:return FV(raw=raw,confidence=confidence,status="needs_review",sources=[source_kind],source_ref=ref,evidence_text=raw)
        normal=format(d.quantize(Decimal("0.0001"),rounding=ROUND_HALF_UP).normalize(),"f")
        return FV(value=normal,raw=raw,confidence=confidence,status="accepted",sources=[source_kind],source_ref=ref,evidence_text=raw)
    if name=="hsn_sac":
        digits=re.fullmatch(r"\d+",str(raw).strip())
        if digits and len(raw) not in {4,6,8}:
            size=next((n for n in (4,6,8) if n>=len(raw)),len(raw));raw=str(raw).zfill(size);normal=raw
            return FV(value=normal,raw=raw,confidence=confidence,status="accepted",sources=[source_kind],source_ref=ref,evidence_text=raw,note="HSN_PADDED")
    if name in NUMERIC_KEYS:
        normal,issues=parse_money(raw)
    elif name in {"invoice_date","due_date"}:normal,issues=_date_value(raw)
    elif name=="description":
        # Preserve intentional line breaks inside quoted CSV/XLSX cell text.
        normal=raw.strip();issues=[]
    else:normal,issues=normalize_field(name,raw)
    return FV(value=normal,raw=raw,confidence=confidence,status="accepted" if not issues else "needs_review",sources=[source_kind],source_ref=ref,evidence_text=raw)

def _layout_b_values(rows,header_row_index):
    found={};upper=rows[:header_row_index]
    label_to_key={_norm(a):k for k,als in ALIASES.items() for a in als}
    for ri,row in enumerate(upper):
        for ci,cell in enumerate(row):
            text=_clean(cell);label=text;inline=None
            if ":" in text:label,inline=text.split(":",1);inline=inline.strip() or None
            key=label_to_key.get(_norm(label))
            if not key:continue
            value=inline
            if not value:
                for j in range(ci+1,min(ci+4,len(row))):
                    if _clean(row[j]):value=_clean(row[j]);break
            if not value and ri+1<len(upper) and ci<len(upper[ri+1]):value=_clean(upper[ri+1][ci]) or None
            if value and key not in found:found[key]=(value,ri+1)
    # Two GSTINs: label proximity classifies seller/customer values.
    for ri,row in enumerate(upper):
        context=" ".join(_clean(v).lower() for v in row)
        for ci,cell in enumerate(row):
            text=_clean(cell);gst=text if re.fullmatch(r"\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]",text.upper()) else None
            if not gst:continue
            target="buyer_gstin" if any(x in context for x in ("buyer","bill to","customer","recipient")) else "supplier_gstin"
            if target not in found:found[target]=(gst,ri+1)
    return found

def _classify(row,headers,mapping,desc_i,last_line):
    cells=[_clean(v) for v in row]; nonempty=[v for v in cells if v]
    if not nonempty:return "blank"
    if len(cells)>=len(headers) and all(not _clean(v) or _norm(v)==_norm(headers[i]) for i,v in enumerate(cells[:len(headers)])):return "header_repeat"
    first=next((v for v in cells if v),"")
    if TOTAL_RE.match(first):return "total"
    numeric_count=sum(_decimal(v) is not None for v in nonempty)
    if last_line and desc_i is not None and (desc_i>=len(cells) or not cells[desc_i]) and numeric_count:return "total"
    if desc_i is not None and len(nonempty)==1 and desc_i<len(cells) and cells[desc_i] and numeric_count==0:
        return "wrapped" if last_line is not None else "note"
    if numeric_count==0 and len(nonempty)==1:return "note"
    return "line_item"

def _make_record(source,invoice_raw,line_raws,refs,sheet,group_id,warnings,table_rows):
    invoice={k:FV() for k in HEADER_FIELDS};lines=[]
    for vals,ref in invoice_raw:
        for key in HEADER_FIELDS:
            raw=vals.get(key)
            if key=="total_amount":raw=vals.get("invoice_total") or raw
            if raw and invoice[key].value is None:invoice[key]=_mk_fv(key,raw,source.input_kind,ref)
    for vals,ref in line_raws:
        line={k:FV() for k in LINE_FIELDS}
        for key in LINE_FIELDS:
            raw=vals.get(key)
            if key=="tax_rate" and not raw:
                rates=[_decimal(vals.get(k)) for k in ("cgst_rate","sgst_rate","igst_rate") if vals.get(k)]
                rates=[v*100 if v<=1 and v*100 in GST_RATES else v for v in rates if v is not None]
                if rates:raw=format(sum(rates).normalize(),"f")
            if raw:line[key]=_mk_fv(key,raw,source.input_kind,ref)
        if any(f.value is not None for f in line.values()):lines.append(line)
    rec=InvoiceRecord(record_id=str(uuid.uuid4()),source=source,invoice=invoice,line_items=lines)
    if not lines:warnings.append(Check(code="NO_LINE_ITEMS",severity="warning",passed=True,message=f"{sheet}: invoice {group_id} has no line-item columns"))
    # Compare declared taxable total against line values using Decimal only.
    if lines:
        for field,components,tolerance in (("taxable_value",("taxable_value",),Decimal("0.01")),("total_amount",("line_total",),Decimal("0.02"))):
            declared=invoice.get(field)
            if not declared or not declared.value:continue
            try:
                component=components[0]
                present=[Decimal(x[component].value) for x in lines if component in x and x[component].value is not None]
                if not present and field=="total_amount":
                    present=[sum((Decimal(x[k].value) for k in ("taxable_value","cgst","sgst","igst","cess") if k in x and x[k].value),Decimal(0)) for x in lines]
                if not present:continue
                computed=sum(present,Decimal(0));diff=abs(Decimal(declared.value)-computed)
                if diff>tolerance:rec.validation.checks.append(Check(code="DECLARED_TOTAL_MISMATCH",severity="warning",passed=False,field_paths=[f"invoice.{field}"],message=f"Declared {field} differs from the sum of line items",expected=format(computed,".2f"),actual=declared.value,difference=format(diff,".2f")))
            except InvalidOperation:pass
    validate(rec,repair=False)
    for out in table_rows:
        for k,f in invoice.items():
            target="invoice_total" if k=="total_amount" else k
            if target in FLAT_COLUMNS:out[target]=f.value or ""
    return rec

def run(path:Path,source:RecordSource)->TabularOutput:
    path=Path(path);warnings=[]
    if source.input_kind=="csv":
        headers,rows,encoding,warnings,preamble=_read_csv(path);tabs=[("Sheet1",headers,rows,0,preamble)]
    elif source.input_kind=="xlsx":
        tabs,warnings=_read_xlsx(path)
    else:raise PipelineError("UNSUPPORTED_FORMAT","Tabular reader accepts only CSV and XLSX")
    records=[];table=[];mappings=[];unmapped=[]
    for sheet,headers,body,header_idx,preamble in tabs:
        rawrows=[row for _,row in body];mapping=_mapping(headers,rawrows)
        for ci,(canonical,method,confidence) in mapping.items():
            flat_target="tax_rate" if canonical in RATE_KEYS else canonical
            mappings.append({"sheet":sheet,"source_column":headers[ci],"canonical_column":flat_target,"method":method,"confidence":confidence})
        for ci,h in enumerate(headers):
            if ci not in mapping:unmapped.append({"sheet":sheet,"source_column":h,"canonical_column":None})
        layout_values=_layout_b_values(preamble,len(preamble)) if preamble else {}
        grouped={};last_key=None;declared_by_group={};declared_rows={}
        desc_i=next((i for i,(k,_,_) in mapping.items() if k=="description"),None)
        for pos,(source_row,row) in enumerate(body):
            vals=_valmap(row,mapping);kind=_classify(row,headers,mapping,desc_i,grouped.get(last_key,{}).get("lines",[]) if last_key else None)
            for source_i,(mapped_key,_,_) in mapping.items():
                raw_id=_clean(row[source_i]) if source_i<len(row) else ""
                if mapped_key in {"invoice_no","supplier_gstin","buyer_gstin","hsn_sac"} and re.fullmatch(r"\d+(?:\.\d+)?[eE][+-]?\d+",raw_id):
                    warnings.append(Check(code="IDENTIFIER_PRECISION_LOSS",severity="warning",passed=True,field_paths=[mapped_key],message=f"Scientific notation retained as text in {sheet}!R{source_row}"))
            if kind in {"blank","header_repeat","note"}:continue
            invoice_no=vals.get("invoice_no") or (last_key[0] if last_key else "") or "__sheet__"
            if not vals.get("invoice_no") and "invoice_no" in {v[0] for v in mapping.values()} and last_key:
                warnings.append(Check(code="INVOICE_NUMBER_INHERITED",severity="warning",passed=True,field_paths=["invoice.invoice_no"],message=f"Invoice number inherited at {sheet}!R{source_row}"))
            supplier=vals.get("supplier_gstin","")
            key=(invoice_no,supplier)
            if "invoice_no" in {v[0] for v in mapping.values()} and invoice_no!="__sheet__":vals["invoice_no"]=invoice_no
            if vals.get("invoice_no"):
                if not supplier and last_key and last_key[0]==invoice_no:supplier=last_key[1];key=(invoice_no,supplier)
                last_key=key
            elif last_key and vals.get("supplier_gstin"):last_key=(invoice_no,supplier);key=last_key
            elif last_key:key=last_key
            else:last_key=key
            if key not in grouped:grouped[key]={"invoice":[],"lines":[],"table":[],"declared":[]}
            ref=f"{sheet}!R{source_row}"
            if kind=="wrapped" and grouped[key]["lines"]:
                prior_vals,prior_ref=grouped[key]["lines"][-1]
                previous=prior_vals.get("description","");wrapped=vals.get("description","")
                prior_vals["description"]=(previous+" "+wrapped).strip()
                if grouped[key]["table"]:
                    grouped[key]["table"][-1]["description"]=(grouped[key]["table"][-1].get("description","")+" "+wrapped).strip()
                continue
            if kind=="total":
                grouped[key]["declared"].append((vals,ref));continue
            # A tabular row's header-level fields are repeated and retained; line fields become a line item.
            grouped[key]["invoice"].append((vals,ref))
            item_columns=any(k in {v[0] for v in mapping.values()} for k in ("description","line_no","hsn_sac","quantity","unit","unit_price","line_total"))
            has_line=item_columns and (any(vals.get(k) for k in LINE_FIELDS if k!="tax_rate") or any(vals.get(k) for k in RATE_KEYS))
            if has_line:grouped[key]["lines"].append((vals,ref))
            flat={c:"" for c in FLAT_COLUMNS}
            flat["source_sheet"]=sheet;flat["source_row"]=source_row
            for k,v in vals.items():
                target="tax_rate" if k in RATE_KEYS else k
                if target in flat:
                    if k in RATE_KEYS:
                        parts=[_decimal(vals.get(rk)) for rk in RATE_KEYS if vals.get(rk)]
                        parts=[x*100 if x<=1 and x*100 in GST_RATES else x for x in parts if x is not None]
                        flat[target]=format(sum(parts).normalize(),"f") if parts else v
                    else:flat[target]=_clean(v)
            extra={headers[i]:_clean(row[i]) if i<len(row) else "" for i in range(len(headers)) if i not in mapping}
            flat["extra"]=extra;table.append(flat);grouped[key]["table"].append(flat)
        for key,data in grouped.items():
            invrows=data["invoice"]
            # Form-like headers come from key/value cells above the line-item header.
            for name,(value,rowno) in layout_values.items():invrows.insert(0,({name:value},f"{sheet}!R{rowno}"))
            # Totals rows provide header totals without being mistaken for line items.
            invrows=data["declared"]+invrows
            rec=_make_record(source,invrows,data["lines"],[x[1] for x in data["lines"]],sheet,key[0],warnings,data["table"])
            records.append(rec)
    return TabularOutput(records=records,table_rows=table,mapping=mappings,unmapped_columns=unmapped,warnings=warnings)
