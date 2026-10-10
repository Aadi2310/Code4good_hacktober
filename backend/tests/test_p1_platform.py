from __future__ import annotations
import io,json,zipfile
from datetime import datetime
from pathlib import Path
from time import sleep
import asyncio
from openpyxl import Workbook,load_workbook
import pytest
from vyom.errors import PipelineError
from vyom.intake import detect
from vyom.models import RecordSource
from vyom.tabular import run
from vyom.outputs.export import csv_bytes,xlsx_bytes,exportable
from vyom.dedupe import find
from vyom.jobs import store
from conftest import make_source

GSTIN="27ABCDE1234F1Z5"

def csv_source(tmp_path,text,encoding="utf-8",name="in.csv"):
    p=tmp_path/name;p.write_bytes(text.encode(encoding));return p,make_source(p)

def run_coroutine_without_io(coro):
    try:coro.send(None)
    except StopIteration as done:return done.value
    raise AssertionError("test coroutine unexpectedly performed asynchronous I/O")

def test_EC01_empty_and_corrupt_inputs(tmp_path):
    p=tmp_path/"zero.csv";p.write_bytes(b"")
    with pytest.raises(PipelineError,match="empty"):detect(p)
    bad=tmp_path/"bad.xlsx";bad.write_bytes(b"PK\x03\x04not a zip")
    with pytest.raises(PipelineError):detect(bad)

def test_EC02_content_wins_over_extension(tmp_path):
    p=tmp_path/"misnamed.csv";p.write_bytes(b"%PDF-1.7\n")
    d=detect(p);assert d.kind=="pdf" and d.ext_mismatch is True

def test_EC03_pdf_signature_routes_without_filename_trust(tmp_path):
    p=tmp_path/"invoice.bin";p.write_bytes(b"%PDF-1.7\n")
    assert detect(p).kind=="pdf"

def test_EC04_xlsx_expansion_limit_is_checked(tmp_path,monkeypatch):
    import vyom.tabular.reader as reader
    # A valid minimal XLSX ZIP whose declared member size exceeds the configured threshold.
    p=tmp_path/"bomb.xlsx"
    with zipfile.ZipFile(p,"w") as z:
        z.writestr("xl/workbook.xml","<workbook/>")
        z.writestr("xl/large.xml",b"x"*2048)
    monkeypatch.setattr(reader,"MAX_UNCOMPRESSED_BYTES",100)
    d=detect(p);assert d.kind=="xlsx"
    with pytest.raises(PipelineError) as err:run(p,make_source(p,"xlsx"))
    assert err.value.code=="ARCHIVE_TOO_LARGE"

def test_EC42_totals_rows_are_not_line_items_and_are_cross_checked(tmp_path):
    text=("Invoice No,Invoice Date,Supplier GSTIN,Description,HSN,Qty,Unit Price,Taxable Value,CGST Rate,SGST Rate,Line Total,Invoice Total,Unknown\n"
          f"INV-1,2026-10-01,{GSTIN},Widget,123,2,50,100,9,9,118,,keep-me\n"
          ",,,,,,,90,,,,90,declared\n")
    p,src=csv_source(tmp_path,text);result=run(p,src)
    assert len(result.records)==1 and len(result.records[0].line_items)==1
    assert result.records[0].validation.checks[0].code=="DECLARED_TOTAL_MISMATCH"
    assert set(result.table_rows[0])==set(__import__("vyom.models",fromlist=["FLAT_COLUMNS"]).FLAT_COLUMNS)
    assert result.table_rows[0]["extra"]["Unknown"]=="keep-me"

def test_EC43_repeated_headers_wrapped_description_and_blank_tail(tmp_path):
    text=("Invoice No,Description,Qty,Unit Price,Taxable Value,Line Total\nINV-1,Long,1,5,5,5\n"
          ",continuation,,,,\nInvoice No,Description,Qty,Unit Price,Taxable Value,Line Total\n"
          "INV-2,Second,1,7,7,7\n\n\n\nINV-3,Should not read,1,9,9,9\n")
    p,src=csv_source(tmp_path,text);result=run(p,src)
    assert len(result.records)==2
    assert result.records[0].line_items[0]["description"].value=="Long continuation"

def test_EC44_hsn_padding_preserves_text_zeroes(tmp_path):
    p,src=csv_source(tmp_path,f"Invoice No,Description,HSN,Qty\nINV-1,Item,00123,1\n")
    rec=run(p,src).records[0]
    assert rec.line_items[0]["hsn_sac"].value=="000123"

def test_EC45_excel_dates_use_workbook_epoch(tmp_path):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    p=tmp_path/"epoch.xlsx";wb=Workbook();ws=wb.active;wb.epoch=CALENDAR_MAC_1904
    ws.append(["Invoice No","Invoice Date","Description","Qty","Unit Price","Taxable Value"])
    ws.append(["INV-1",datetime(2024,1,2),"item",1,2,2]);wb.save(p)
    rec=run(p,make_source(p,"xlsx")).records[0]
    assert rec.invoice["invoice_date"].value=="2024-01-02"

def test_EC46_windows_1252_semicolon_and_quoted_newline(tmp_path):
    text=f"Invoice No;Supplier;Supplier GSTIN;Description;Qty;Taxable Value\nINV-1;Café;{GSTIN};\"one\ntwo\";1;2\n"
    p,src=csv_source(tmp_path,text,"cp1252");result=run(p,src)
    assert len(result.records)==1 and result.records[0].line_items[0]["description"].value=="one\ntwo"
    assert result.records[0].line_items[0]["description"].source_ref=="Sheet1!R2"

def test_EC47_rate_forms_are_normalized(tmp_path):
    p,src=csv_source(tmp_path,"Invoice No,Description,Taxable Value,CGST Rate,SGST Rate\nINV-1,Item,100,0.09,9%\n")
    line=run(p,src).records[0].line_items[0]
    assert line["tax_rate"].value=="18"

def test_layout_B_key_value_header_cells_are_attached(tmp_path):
    text=("Invoice No,INV-B,,Supplier,Acme\nInvoice Date,2026-05-01,,Buyer,Beta\n"
          "Description,Qty,Taxable Value\nWidget,2,20\n")
    p,src=csv_source(tmp_path,text);rec=run(p,src).records[0]
    assert rec.invoice["invoice_no"].value=="INV-B" and rec.invoice["supplier_name"].value=="Acme"
    assert rec.line_items[0]["description"].value=="Widget"

def test_layout_C_gstr_style_keeps_no_line_items(tmp_path):
    text=(f"Invoice No,Invoice Date,Supplier GSTIN,Taxable Value,Total Tax,Invoice Total\n"
          f"INV-C,2026-05-01,{GSTIN},100,18,118\nINV-C,2026-05-01,{GSTIN},50,6,56\n")
    p,src=csv_source(tmp_path,text);result=run(p,src)
    assert len(result.records)==1 and result.records[0].line_items==[]
    assert len(result.table_rows)==2 and result.records[0].invoice["total_amount"].value=="118.00"

def test_two_level_tax_headers_map_to_total_rate(tmp_path):
    p=tmp_path/"two-level.csv"
    p.write_text("Invoice No,Description,Taxable Value,CGST,SGST\n,, ,Rate,Rate\nINV-1,item,100,9,9\n",encoding="utf-8")
    line=run(p,make_source(p)).records[0].line_items[0]
    assert line["tax_rate"].value=="18"

def test_EC48_hidden_sheets_rows_and_columns_are_ignored(tmp_path):
    p=tmp_path/"hidden.xlsx";wb=Workbook();ws=wb.active;ws.title="Visible"
    ws.append(["Invoice No","Description","Qty","Taxable Value"]);ws.append(["INV-1","ok",1,1]);ws.append(["INV-2","hidden",1,1]);ws.row_dimensions[3].hidden=True;ws.column_dimensions["D"].hidden=True
    hidden=wb.create_sheet("Hidden");hidden.sheet_state="hidden";hidden.append(["Invoice No"]);hidden.append(["INV-X"]);wb.save(p)
    result=run(p,make_source(p,"xlsx"))
    assert len(result.records)==1 and any(w.code=="HIDDEN_CONTENT_IGNORED" for w in result.warnings)

def test_EC49_uncalculated_formula_is_reported(tmp_path):
    p=tmp_path/"formula.xlsx";wb=Workbook();ws=wb.active;ws.append(["Invoice No","Description","Qty","Taxable Value"]);ws.append(["INV-1","item",1,"=SUM(1,2)"]);wb.save(p)
    result=run(p,make_source(p,"xlsx"));assert any(w.code=="FORMULAS_NOT_CALCULATED" for w in result.warnings)

def test_EC50_unmapped_columns_are_preserved(tmp_path):
    p,src=csv_source(tmp_path,"Invoice No,Description,Qty,Secret Ref\nINV-1,item,1,retain-this\n")
    result=run(p,src);assert result.unmapped_columns and result.table_rows[0]["extra"]["Secret Ref"]=="retain-this"

def test_EC53_exact_duplicate_upload_is_cached(tmp_path):
    from vyom.api import main
    content=f"Invoice No,Description,Qty,Taxable Value\nINV-1,item,1,1\n".encode()
    p=tmp_path/"inv.csv";p.write_bytes(content);detected=detect(p);sha=__import__("hashlib").sha256(content).hexdigest()
    first=main._queue(p,"inv.csv",sha,detected,"batch-1",{"force":False})
    for _ in range(100):
        if store.get(first["job_id"])["status"] in {"completed","failed"}:break
        sleep(.02)
    assert store.get(first["job_id"])["status"]=="completed"
    second=main._queue(p,"inv.csv",sha,detected,"batch-2",{"force":False})
    assert second["cached"] is True and second["job_id"]==first["job_id"]

def test_EC54_interrupted_worker_requeues_once(tmp_path,monkeypatch):
    monkeypatch.setattr(store,"DATA_DIR",tmp_path);monkeypatch.setattr(store,"DB_PATH",tmp_path/"db.sqlite")
    job,_=store.create("x.csv","a"*64,"csv")
    store.update(job["id"],"processing",stage="parsing")
    recovered=store.recover_running();assert len(recovered)==1 and store.get(job["id"])["status"]=="queued"

def test_EC55_timeout_keeps_partial_output(monkeypatch,tmp_path):
    from vyom.api import main
    from vyom.models import RecordSource,TabularOutput
    source=RecordSource(filename="x.csv",sha256="b"*64,detected_type="text/csv",input_kind="csv")
    job,_=store.create("x.csv",source.sha256,"csv")
    artifact=tmp_path/"artifacts"/job["id"];artifact.mkdir(parents=True);p=artifact/"source.bin";p.write_text("x")
    monkeypatch.setattr(main,"tabular_run",lambda *args:TabularOutput(records=[],table_rows=[],mapping=[]))
    ticks=iter([0.,2.,3.,4.]);monkeypatch.setattr(main.time,"perf_counter",lambda:next(ticks,4.))
    main._process(job["id"],p,source,{"timeout_s":1})
    assert store.get(job["id"])["status"]=="failed" and store.result(job["id"])["documents"]==[]

def test_EC57_csv_and_xlsx_exports_escape_formula_injection(tmp_path):
    p,src=csv_source(tmp_path,"Invoice No,Description,Qty\nINV-1,=HYPERLINK(\"x\"),1\n")
    records=run(p,src).records
    data=csv_bytes(records);assert b"'=HYPERLINK" in data
    book=load_workbook(io.BytesIO(xlsx_bytes(records)),read_only=True,data_only=True)
    assert book.sheetnames==["Invoices","LineItems","Validation","Corrections","Audit"]
    assert book["LineItems"]["I2"].value=="'=HYPERLINK(\"x\")"

def test_review_edit_approve_and_export_policy(tmp_path):
    from vyom.api import main
    from vyom.api.main import PatchRequest,Edit,ReviewRequest
    p,source=csv_source(tmp_path,"Invoice No,Description,Qty,Taxable Value,Invoice Total\nINV-1,item,1,10,10\n")
    rec=run(p,source).records[0]
    job,_=store.create("review.csv",source.sha256,"csv",force=True);jid=job["id"]
    envelope={"job_id":jid,"status":"completed","input":source.model_dump(mode="json"),"documents":[rec.model_dump(mode="json")],"tabular":None,"error":None,"artifacts":{"pages":[]},"timings_ms":{}}
    store.update(jid,"completed",envelope)
    edit=run_coroutine_without_io(main.patch_doc(jid,0,PatchRequest(actor="reviewer",edits=[Edit(path="invoice.invoice_no",value="INV-2")])))
    assert edit["invoice"]["invoice_no"]["status"]=="human_verified"
    reviewed=run_coroutine_without_io(main.review(jid,0,ReviewRequest(action="approve",actor="reviewer")))
    assert reviewed["review"]["state"]=="approved"
    exported=main.export(jid,"xlsx","all");assert exported.status_code==200
    audit=load_workbook(io.BytesIO(exported.body),read_only=True)["Audit"]
    assert audit.max_row==3

def test_mock_result_fixture_is_ui_usable():
    from vyom.api.main import mock_result
    data=mock_result()
    assert len(data["documents"][0]["line_items"])==2
    assert data["documents"][0]["review"]["required"] is True

def test_dedupe_business_and_near_match(tmp_path,monkeypatch):
    from vyom.models import InvoiceRecord,FV,Quality,Review,Validation
    from vyom.models import RecordSource as RS
    source=RS(filename="a.csv",sha256="a"*64,detected_type="text/csv",input_kind="csv")
    fields={k:FV() for k in __import__("vyom.models",fromlist=["HEADER_FIELDS"]).HEADER_FIELDS}
    fields["supplier_gstin"]=FV(value=GSTIN);fields["invoice_no"]=FV(value="INV-001");fields["invoice_date"]=FV(value="2026-01-01");fields["total_amount"]=FV(value="118.00");fields["supplier_name"]=FV(value="ACME LTD")
    record=InvoiceRecord(record_id="r1",source=source,invoice=fields)
    matches=find(record,None,other_records=[("prior",0,record.model_copy(update={"record_id":"r2","source":source.model_copy(update={"sha256":"c"*64})}))])
    assert matches and matches[0].kind=="business_key"
